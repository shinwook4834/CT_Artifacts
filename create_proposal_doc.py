import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_left_border(cell, color_hex="1B365D", sz="36"):
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def set_table_borders(table, color="D2D6DC", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def build_document():
    doc = docx.Document()

    # Set Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Style definitions
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Malgun Gothic'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x2D, 0x37, 0x48) # Slate charcoal
    normal_style.paragraph_format.line_spacing = 1.35
    normal_style.paragraph_format.space_after = Pt(4)

    # Palette
    NAVY = RGBColor(0x1B, 0x36, 0x5D)
    STEEL_BLUE = RGBColor(0x2B, 0x6C, 0xB0)
    DARK_GRAY = RGBColor(0x4A, 0x55, 0x68)
    MUTED_RED = RGBColor(0x9B, 0x2C, 0x2C)

    # --- Title Block ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_p.paragraph_format.space_after = Pt(2)
    run_tag = title_p.add_run("[연구 계획서 / Research Proposal]\n")
    run_tag.font.name = 'Malgun Gothic'
    run_tag.font.size = Pt(11)
    run_tag.font.bold = True
    run_tag.font.color.rgb = STEEL_BLUE

    run_title = title_p.add_run("기치료 유방암 환자의 실제 치료계획 기반 케모포트 AI 인공음영 저감(MAR) 및 후향적 선량 재평가 연구")
    run_title.font.name = 'Malgun Gothic'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    eng_title_p = doc.add_paragraph()
    eng_title_p.paragraph_format.space_after = Pt(14)
    run_eng = eng_title_p.add_run("Retrospective Dosimetric Evaluation of Breast Cancer Radiation Therapy with a Novel CAD Prior-Guided AI Metal Artifact Reduction for Patients with Chemo Ports")
    run_eng.font.name = 'Calibri'
    run_eng.font.size = Pt(11.5)
    run_eng.font.italic = True
    run_eng.font.color.rgb = DARK_GRAY

    # --- Metadata Table ---
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table, color="CBD5E0", sz="6")
    meta_data = [
        ("연구 유형", "후향적 임상-의학물리 융합 연구 (Retrospective Clinical & Dosimetric Study)"),
        ("목표 저널군", "Medical Physics, Radiotherapy and Oncology (Green Journal), PMB, IJROBP (Red Journal)"),
        ("연구 대상", "케모포트(Chemo port) 삽입 상태로 쇄골상부 림프절/흉벽 방사선 치료를 완료한 유방암 환자 코호트")
    ]
    col_widths = [Inches(1.5), Inches(5.0)]
    for row_idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = col_widths[0], col_widths[1]
        set_cell_background(c0, "EDF2F7")
        set_cell_background(c1, "F7FAFC")
        set_cell_margins(c0, top=100, bottom=100, left=140, right=140)
        set_cell_margins(c1, top=100, bottom=100, left=140, right=140)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = NAVY
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- Section 1: 연구 배경 및 핵심 가설 ---
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. 연구 배경 및 핵심 가설 (Background & Hypotheses)")
    r_h1.font.name = 'Malgun Gothic'
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = NAVY
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "• 임상적 현주소: 전 세계 유방암 환자의 항암치료에 사용되는 케모포트의 약 60%는 여전히 고밀도 티타늄(Titanium) 재질입니다. "
        "선행 및 수술 후 방사선 치료 시 케모포트가 쇄골상부 림프절(SCN) 및 흉벽 치료 부위에 위치하면서 심각한 금속 인공음영(Metal Artifact)을 유발합니다."
    )
    doc.add_paragraph(
        "• 기존 임상 치료의 한계: 상용 MAR(iMAR, O-MAR)은 케모포트처럼 '피부 바로 밑(Subcutaneous)'이자 '폐(공기) 경계면'에 인접한 부위에서 2차 암영과 블러링을 발생시킵니다. "
        "또한 포트 내부의 비금속 물질(실리콘 셉텀, 유체 챔버)을 무시하고 전체를 통짜 금속 또는 물로 강제 지정(Bulk Override)하여 치료계획을 수립해 왔습니다."
    )

    # Callout Box: Core Clinical Question
    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = callout.rows[0].cells[0]
    c.width = Inches(6.5)
    set_cell_background(c, "F0F4F8")
    set_cell_left_border(c, color_hex="1B365D", sz="36")
    set_cell_margins(c, top=120, bottom=120, left=180, right=140)
    cp = c.paragraphs[0]
    cp.paragraph_format.line_spacing = 1.3
    c_bold = cp.add_run("핵심 임상 질문 (Clinical Question):\n")
    c_bold.font.bold = True
    c_bold.font.size = Pt(10)
    c_bold.font.color.rgb = NAVY
    c_text = cp.add_run(
        "“그렇다면, 우리가 과거에 이미 치료를 마친 환자들의 실제 몸속에는 선량이 얼마나 정확하게 들어갔던 것일까? "
        "아티팩트 때문에 의사가 타깃(CTV)을 왜곡되게 그렸거나, 포트 뒤쪽 림프절에 심각한 선량 결손(Cold spot)이 발생하지는 않았는가?”"
    )
    c_text.font.size = Pt(10)
    c_text.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Hypothesis
    p_hypo = doc.add_paragraph()
    r_hypo_title = p_hypo.add_run("• 연구 가설 (Hypothesis): ")
    r_hypo_title.font.bold = True
    r_hypo_title.font.color.rgb = STEEL_BLUE
    p_hypo.add_run(
        "제조사 규격 3D CAD Prior와 2단계 AI(자세추정 + 조건부 인페인팅)를 통해 '진짜 해부학 및 포트 내부 밀도'가 복원된 CT 위에, "
        "기존에 승인된 실제 치료 빔(Beam, MU, Segment)을 그대로 재계산(Frozen-Plan Recalculation)하면, "
        "기존 치료계획 시스템(TPS)이 예측했던 선량과 임상적으로 유의미한 표적 선량 결손(Cold spot) 및 피부 과선량(Hot spot)이 실증될 것이다."
    )

    # --- Section 2: 연구 재료 및 방법 ---
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. 연구 재료 및 방법 (Materials and Methods)")
    r_h2.font.name = 'Malgun Gothic'
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = NAVY
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    # 2.1
    p_sub1 = doc.add_paragraph()
    r_sub1 = p_sub1.add_run("2.1 환자 코호트 및 수집 데이터 (Retrospective Cohort)")
    r_sub1.font.bold = True
    r_sub1.font.size = Pt(11.5)
    r_sub1.font.color.rgb = STEEL_BLUE
    doc.add_paragraph(
        "• 대상군: 케모포트를 유지한 채 쇄골상부 림프절(SCN) 및 흉벽/유방을 포함하여 방사선 치료(VMAT 또는 3D-CRT/IMRT)를 완료한 유방암 환자 N명(20~30례).\n"
        "• 수집 DICOM-RT 데이터: 원본 모의치료 CT, 상용 MAR 적용 CT, 승인된 치료계획(RP/Plan), 방사선종양학과 의사가 묘사한 구조물(RS/Structure), 계산된 3차원 선량 분포(RD/Dose)."
    )

    # 2.2
    p_sub2 = doc.add_paragraph()
    r_sub2 = p_sub2.add_run("2.2 제조사 레퍼런스 규격 기반 다중 물질 3D CAD Prior 구축")
    r_sub2.font.bold = True
    r_sub2.font.size = Pt(11.5)
    r_sub2.font.color.rgb = STEEL_BLUE
    doc.add_paragraph(
        "환자에게 삽입된 케모포트(예: BD PowerPort® Titanium)의 공식 기술 데이터시트(Specification Sheet)를 바탕으로, "
        "통짜 금속이 아닌 3가지 개별 물리 구획으로 모델링된 파라메트릭 3D CAD를 구축합니다:"
    )

    # Material Table
    mat_table = doc.add_table(rows=4, cols=3)
    mat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(mat_table, color="CBD5E0", sz="4")
    headers = ["구분 (Compartment)", "물리적 재질 및 전자밀도 (Density)", "방사선 치료학적 의미"]
    for i, h in enumerate(headers):
        cell = mat_table.rows[0].cells[i]
        set_cell_background(cell, "1B365D")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    mat_rows = [
        ("Part 1: 외벽 및 바닥 챔버", "Titanium Alloy (ρ = 4.43 g/cm³, HU > 7000)", "방사선 차폐(Cold spot) 및 경계면 후방산란 유발"),
        ("Part 2: 실리콘 격막 (Septum)", "Silicone Elastomer (ρ = 1.15 g/cm³, HU ~ 120)", "연부조직과 유사한 투과 특성 (금속으로 오판 방지)"),
        ("Part 3: 내부 유체 챔버", "Fluid Cavity (ρ = 1.00 g/cm³, HU ~ 0)", "식염수 충만 상태 모사 (광자선이 자유롭게 통과)")
    ]
    for r_idx, r_data in enumerate(mat_rows):
        row = mat_table.rows[r_idx + 1]
        for c_idx, text in enumerate(r_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, "F7FAFC" if r_idx % 2 == 0 else "EDF2F7")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9.0)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 2.3
    p_sub3 = doc.add_paragraph()
    r_sub3 = p_sub3.add_run("2.3 2-Stage Prior-Guided AI 복원 프레임워크")
    r_sub3.font.bold = True
    r_sub3.font.size = Pt(11.5)
    r_sub3.font.color.rgb = STEEL_BLUE
    doc.add_paragraph(
        "• Stage 1 (6-DoF 자세 추정 및 포트 내부 고정): "
        "가벼운 3D CNN 기반 네트워크가 아티팩트로 번진 CT 속에서 케모포트의 3차원 위치(x, y, z)와 회전각(Roll, Pitch, Yaw)을 정확히 추정합니다. "
        "규격화된 3D CAD 모델을 완벽히 중첩 정합(Rigid Registration)하여, 포트 외벽은 진짜 크기로 확정하고 내부는 실제 티타늄/실리콘/유체 밀도로 강제 고정합니다.\n"
        "• Stage 2 (조건부 인페인팅 Diffusion/Transformer): "
        "포트 외벽(Inner condition)과 줄무늬 밖의 정상 인체 해부학(Outer condition: 피부선, 늑골, 정상 폐 실질)을 조건으로 부여합니다. "
        "Conditional Diffusion Model(RePaint/Palette 구조)을 통해 줄무늬를 제거하고, 상용 MAR이 뭉개뜨리던 쇄골하혈관, 림프절 구획(SCN), 피부 경계선을 완벽히 복원합니다."
    )

    # 2.4
    p_sub4 = doc.add_paragraph()
    r_sub4 = p_sub4.add_run("2.4 동결 플랜 선량 재계산 방법론 (Frozen-Plan Recalculation)")
    r_sub4.font.bold = True
    r_sub4.font.size = Pt(11.5)
    r_sub4.font.color.rgb = STEEL_BLUE
    doc.add_paragraph(
        "임상적 핵심 기법: 환자의 기존 치료계획에 포함된 모든 빔 파라미터(빔 각도, 세그먼트 형태, 모니터 유닛[MU], 갠트리 궤적 등)를 100% 동결(Freeze)합니다.\n"
        "동일한 치료 빔을 다음 3개 CT 세트에 조사하여 선량만 순방향 재계산(Forward Recalculation)합니다:\n"
        "  1. Plan_Orig : 아티팩트가 있는 기존 오리지널 CT 기반 선량 (치료 당시 승인 선량)\n"
        "  2. Plan_Comm : 상용 MAR(iMAR 등) 적용 CT 기반 선량\n"
        "  3. Plan_AI   : 제안한 다중 물질 CAD Prior AI 복원 CT 기반 선량 (실제 환자 몸에 전달된 참 선량)"
    )

    # --- Section 3: 정량적 평가 지표 ---
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. 정량적 평가 지표 (Quantitative Evaluation Endpoints)")
    r_h3.font.name = 'Malgun Gothic'
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = NAVY
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    # Two-track explanation
    p_tt = doc.add_paragraph()
    r_tt_title = p_tt.add_run("3.1 표적 윤곽 변화 및 선량 평가 (Two-Track Target Evaluation)\n")
    r_tt_title.font.bold = True
    r_tt_title.font.size = Pt(11.5)
    r_tt_title.font.color.rgb = STEEL_BLUE
    p_tt.add_run(
        "아티팩트로 인해 혈관과 해부학적 경계가 가려져 치료 당시 타깃(CTV/PTV) 자체가 왜곡되었을 가능성을 고려하여, "
        "두 가지 트랙으로 분리하여 정밀 평가합니다:"
    )

    # Track Table
    track_table = doc.add_table(rows=3, cols=2)
    track_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(track_table, color="CBD5E0", sz="4")
    
    t_headers = ["평가 트랙 (Evaluation Track)", "주요 분석 내용 및 정량 지표"]
    for i, h in enumerate(t_headers):
        cell = track_table.rows[0].cells[i]
        set_cell_background(cell, "1B365D")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    t_rows = [
        ("[Track A]\n기하학적 표적 윤곽 오차\n(Contouring Geometric Error)",
         "• 방법: 복원된 깨끗한 AI-CT에서 전문의가 ESTRO/RTOG 가이드라인에 따라 '진짜 표적(CTV_True, PTV_True)'을 재설정(Re-contouring)\n"
         "• 지표: 체적 차이(ΔV), 일치도(Dice Similarity Coefficient, DSC), 95% Hausdorff Distance (HD95)\n"
         "• 임상 의미: 아티팩트로 인해 치료 당시 의사가 타깃을 얼마나 과소/과대 설정했는가 입증"),
        ("[Track B]\n선량 평가: 물리적 오차 vs\n임상적 표적 누락(Geographic Miss)",
         "• 1) 순수 물리적 선량 오차 (on Frozen CTV_Orig): 타깃을 그대로 뒀을 때 오직 밀도 왜곡 때문에 발생한 선량 오차 (D95%, Dmean, Cold spot 크기)\n"
         "• 2) 임상적 표적 누락 분석 (on Re-contoured CTV_True) ★핵심:\n"
         "  “기존 치료계획(Plan_Orig)은 새로 밝혀진 '진짜 림프절 표적(CTV_True)'을 얼마나 놓치고 있었는가?”\n"
         "  → V95%(CTV_True), D98%, PTV coverage 저하율 산출")
    ]
    for r_idx, r_data in enumerate(t_rows):
        row = track_table.rows[r_idx + 1]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(2.2), Inches(4.3)
        set_cell_background(c0, "EDF2F7")
        set_cell_background(c1, "F7FAFC")
        set_cell_margins(c0, top=100, bottom=100, left=120, right=120)
        set_cell_margins(c1, top=100, bottom=100, left=120, right=120)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(r_data[0])
        r0.font.bold = True
        r0.font.size = Pt(9.0)
        r0.font.color.rgb = NAVY
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(r_data[1])
        r1.font.size = Pt(9.0)
        r1.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 3.2 OAR & Skin
    p_oar = doc.add_paragraph()
    r_oar_title = p_oar.add_run("3.2 정상장기(OAR) 및 피부 선량 평가\n")
    r_oar_title.font.bold = True
    r_oar_title.font.size = Pt(11.5)
    r_oar_title.font.color.rgb = STEEL_BLUE
    p_oar.add_run(
        "• 피부 선량 (Skin Dose): 포트 직상방 피부의 최대선량(Dmax) 및 평균선량(Dmean). 후방산란에 의한 실제 피부 과선량 평가.\n"
        "• 동측 폐 (Ipsilateral Lung): V20Gy(20Gy 이상 받는 폐 체적 비율), V5Gy. 상용 MAR의 계면 블러링으로 인한 폐 선량 과대/과소평가 오차 규명.\n"
        "• 3차원 선량 일치도: 3D Gamma Pass Rate (3%/3mm, 2%/2mm, 1%/1mm 기준)를 통한 Plan_Orig vs Plan_AI 전역 선량 분포 비교."
    )

    # --- Section 4: 예상 결과 및 임상적 기대효과 ---
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. 예상 결과 및 임상적 기대효과 (Expected Impact)")
    r_h4.font.name = 'Malgun Gothic'
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = NAVY
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "1. 미지의 임상적 사실 규명: 과거 치료계획 시스템이 케모포트 아티팩트 및 비금속 내부 물질 무시로 인해 림프절 표적에 얼마만큼의 선량 결손을 유발했는지 최초로 정량 보고.\n"
        "2. 표적 누락(Geographic Miss)의 위험성 경고: 아티팩트로 인해 의사가 타깃을 불완전하게 설정함으로써 발생했던 잠재적 종양 미조사(Under-dosage) 위험을 입증하여 진료 지침 개선에 기여.\n"
        "3. 상용 MAR의 한계 극복 및 워크플로우 혁신: 수동 밀도 오버라이드나 부정확한 상용 알고리즘에 의존하지 않고, AI 기반의 자동화된 초정밀 CT 복원 및 선량 계산 표준 프로토콜 제시."
    )

    # --- Section 5: 방사선 치료 아티팩트 극복 5대 핵심 로드맵 ---
    h5 = doc.add_heading(level=1)
    r_h5 = h5.add_run("5. 향후 연계: 방사선 치료 금속 아티팩트 극복 5대 핵심 로드맵 (The 5-Stage Grand Vision)")
    r_h5.font.name = 'Malgun Gothic'
    r_h5.font.size = Pt(14)
    r_h5.font.bold = True
    r_h5.font.color.rgb = NAVY
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "본 케모포트 연구는 인체 내 고밀도 의료기기로 인해 발생하는 방사선 치료 오차를 완전 정복하기 위한 5부작 연속 연구의 첫 번째 출발점입니다:"
    )

    # Roadmap Table
    rd_table = doc.add_table(rows=6, cols=3)
    rd_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(rd_table, color="CBD5E0", sz="4")
    
    r_headers = ["단계 (Series)", "대상 임플란트 및 부위", "주요 위협 장기 및 핵심 AI 해결책"]
    for i, h in enumerate(r_headers):
        cell = rd_table.rows[0].cells[i]
        set_cell_background(cell, "1B365D")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    r_rows = [
        ("제1탄 (본 연구)", "케모포트 (Chemo Port)\n• 유방암 / 쇄골상부 림프절", "피하층/폐 계면 왜곡 및 다중 물질 CAD Prior\n→ 후향적 선량 재계산 및 표적 누락(Geographic Miss) 실증"),
        ("제2탄 ★신규 확장", "척추경 나사 (Spine Screws)\n• 척추 전이암 SBRT", "1~2mm 인접 척수(Spinal cord) 마비 영구 장애 위험\n→ 나사산 CAD Prior + 척수 PRV 초정밀 복원 및 급격한 선량 감쇄 최적화"),
        ("제3탄 ★신규 확장", "인공 심박동기 (Pacemaker/ICD)\n• 폐암 / 좌측 유방암", "2 Gy 피폭 한계 반도체 기기 소손 및 오작동 사망 방지\n→ 배터리 팩 CAD 정합 + 기기 피폭 선량 정밀화 및 폐종양 선량 확보"),
        ("제4탄", "치아 보철물 (Dental Implants)\n• 두경부암 / 구강·인후두", "다발성 금속체 상호 교차 줄무늬로 침샘/연하근 기능 파괴\n→ 다중 클러스터 Dual-Domain AI 및 정상 침샘/하악골 해부학 복원"),
        ("제5탄", "인공 고관절 (Hip Prosthesis)\n• 골반암 / 전립선·부인과", "초대형 완전 광자 결핍(Total Starvation)에 의한 골반 선량 붕괴\n→ 대형 관절 CAD Prior + 몬테카를로 물리 통합 AI 및 방광/직장 정상화")
    ]
    for r_idx, r_data in enumerate(r_rows):
        row = rd_table.rows[r_idx + 1]
        for c_idx, text in enumerate(r_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, "F7FAFC" if r_idx % 2 == 0 else "EDF2F7")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8.8)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = NAVY

    doc.save("/Users/shinwookkim/Documents/Geant4/ChemoPort_MAR_Research_Proposal.docx")
    print("Successfully created /Users/shinwookkim/Documents/Geant4/ChemoPort_MAR_Research_Proposal.docx")

if __name__ == "__main__":
    build_document()
